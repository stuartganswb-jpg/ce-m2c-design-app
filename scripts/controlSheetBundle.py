#!/usr/bin/env python3
"""Turn the Excel control sheets in Inception/ into files the app's 1.2 Control Sheets tab can import.

    python3 scripts/controlSheetBundle.py

Why a file and not a direct load: the database refuses anything that is not the signed-in app, so the rows and
pictures are prepared here and the IMPORT button on tab 1.2 writes them. One .control-sheet.json per workbook,
pictures inside it, written to Inception/import/. Nothing here touches the app or the database. Needs openpyxl and
Pillow. The file format is read by Shared/controlSheet.js (readBundle / planImport).

Workbooks read: M2CCapa Floor Lamp.xlsx, M2C_Strata Light.xlsx, H3-Contours.xlsx. H1 is NOT read (Stuart 2026-10-07:
"do not import the H1 it is already in the app").
"""
import base64, io, json, os, re, sys, zipfile, datetime, warnings
import xml.etree.ElementTree as ET
import openpyxl
from PIL import Image

warnings.filterwarnings('ignore')
SRC = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'Inception')   # the workbooks (not in git)
HERE = os.path.join(SRC, 'import')                                                                # the files made (not in git)
PART_PX, DRAWING_PX = 520, 2200

NS = {
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'pr': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'rd': 'http://schemas.microsoft.com/office/spreadsheetml/2017/richdata',
    'rvrel': 'http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel',
}


def clean(v):
    if v is None: return ''
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    s = str(v).replace('\r', '').strip()
    return '' if s == '#VALUE!' else s


def num(v):
    """A number, or None. Never guesses: 'CNY 480.00' reads 480, '100m' is not a number."""
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace(',', '').replace('。', '.')
    m = re.fullmatch(r'(?:CNY|RMB|USD|\$)?\s*(-?\d+(?:\.\d+)?)\s*', s, re.I)
    return float(m.group(1)) if m else None


class Pictures:
    """Every picture of one workbook, shrunk once and named — the bundle carries each a single time."""
    def __init__(self):
        self.images, self._by_hash = {}, {}

    def add(self, raw, px, prefix):
        key = hash(raw)
        if key in self._by_hash: return self._by_hash[key]
        im = Image.open(io.BytesIO(raw))
        im.load()
        if max(im.size) > px:
            im.thumbnail((px, px), Image.LANCZOS)
        if im.mode not in ('RGB', 'RGBA', 'L', 'LA', 'P'): im = im.convert('RGBA')
        out = io.BytesIO()
        im.save(out, 'PNG', optimize=True)
        name = f'{prefix}{len(self.images) + 1}'
        self.images[name] = {'mime': 'image/png', 'w': im.size[0], 'h': im.size[1], 'data': base64.b64encode(out.getvalue()).decode('ascii')}
        self._by_hash[key] = name
        return name


def cell_pictures(path):
    """{ sheet title: { 'H3': bytes } } — the pictures Excel keeps INSIDE cells (they read #VALUE! to anything else)."""
    z = zipfile.ZipFile(path)
    names = set(z.namelist())
    if 'xl/richData/rdrichvalue.xml' not in names: return {}
    rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read('xl/richData/_rels/richValueRel.xml.rels'))}
    rel_ids = [r.get('{%s}id' % NS['r']) for r in ET.fromstring(z.read('xl/richData/richValueRel.xml'))]
    rich = [int(rv.find('rd:v', NS).text) for rv in ET.fromstring(z.read('xl/richData/rdrichvalue.xml')).findall('rd:rv', NS)]
    meta = ET.fromstring(z.read('xl/metadata.xml'))
    future = [int(bk.find('.//rd:rvb', NS).get('i')) for bk in meta.find('m:futureMetadata', NS).findall('m:bk', NS)]
    value_meta = [int(bk.find('m:rc', NS).get('v')) for bk in meta.find('m:valueMetadata', NS).findall('m:bk', NS)]
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    wb_rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    out = {}
    for sh in wb.find('m:sheets', NS):
        target = wb_rels[sh.get('{%s}id' % NS['r'])]
        xml = z.read('xl/' + target.lstrip('/').replace('xl/', '')).decode('utf8')
        cells = {}
        for ref, vm in re.findall(r'<c r="([A-Z]+\d+)"[^>]*?\bvm="(\d+)"', xml):
            rv = future[value_meta[int(vm) - 1]]
            media = rels[rel_ids[rich[rv]]].replace('../', 'xl/')
            cells[ref] = z.read(media)
        out[sh.get('name')] = cells
    return out


def floating_pictures(ws):
    """[(row, col, bytes)] — pictures laid over the sheet (the drawings; in H3, the part pictures)."""
    out = []
    for img in getattr(ws, '_images', []):
        fr = getattr(img.anchor, '_from', None)
        out.append((fr.row + 1 if fr else 0, fr.col + 1 if fr else 0, img._data()))
    return out


def header_map(ws, row):
    return {clean(c.value).upper().replace('\n', ' '): c.column for c in ws[row] if clean(c.value)}


def lighting(path, project_name, suggested_brand):
    wb = openpyxl.load_workbook(path, data_only=True)
    in_cell = cell_pictures(path)
    pics, report = Pictures(), []
    sections, lines, by_key = [], [], {}
    title_to_section = {}

    for si, ws in enumerate(wb.worksheets):
        H = header_map(ws, 2)
        col = lambda *names: next((H[n] for n in names if n in H), None)
        c_no, c_id, c_color, c_desc = col('NO#', '#'), col('M2C ID ?'), col('M2C COLOR'), col('DESCRIPTION')
        c_size, c_mat, c_wt, c_pic = col('SIZE'), col('CONTENT'), col('WEIGHT'), col('PICTURE')
        c_rmb, c_qq, c_oq, c_usd, c_duty = col('PRICE/PC RMB'), col('QTY QUOTED'), col('ORDER QTY'), col('USD'), col('DTY/ SHIPPING')
        c_each, c_qty, c_sub = col('TOTAL COST EACH'), col('QTY PER LIGHT'), col('SUBTOTAL', 'COST PER ASSEMBLY')
        c_sku, c_vendor, c_origin, c_status, c_note = col('VENDOR ID'), col('VENDOR'), col('ORIGIN'), col('STATUS'), col('NOTE')
        sec_name = clean(ws.cell(2, 1).value) or clean(ws.cell(1, 1).value) or ws.title
        if len(wb.worksheets) > 1: sec_name = ws.title
        sec = {'key': f'S{si + 1}', 'name': sec_name, 'kind': 'ASSEMBLY', 'drawings': [], 'children': [], 'sheetTotal': None, 'sourceSheet': ws.title}
        for (_r, _c, raw) in sorted(floating_pictures(ws), key=lambda t: -len(t[2])):
            sec['drawings'].append(pics.add(raw, DRAWING_PX, 'd'))
        sections.append(sec)
        title_to_section[ws.title.strip().upper()] = sec

        g = lambda r, c: ws.cell(r, c).value if c else None
        for r in range(3, ws.max_row + 1):
            desc = clean(g(r, c_desc))
            if not desc:
                tot = num(g(r, c_sub))
                if tot is not None and not clean(g(r, c_no)): sec['sheetTotal'] = tot
                continue
            balloon, qty = clean(g(r, c_no)), num(g(r, c_qty))
            child = title_to_section.get(desc.upper())
            if child is not None and child is not sec:          # a line that IS another sheet of this workbook
                sec['kind'] = 'PRODUCT'
                sec['children'].append({'section': child['key'], 'qty': qty if qty is not None else 1, 'balloon': balloon})
                continue
            sku = clean(g(r, c_sku))
            v, o = clean(g(r, c_vendor)), clean(g(r, c_origin))
            url = v if v.lower().startswith('http') else (o if o.lower().startswith('http') else '')
            if v.lower().startswith('http'): vendor, origin = o, ''
            else: vendor, origin = v, ('' if o.lower().startswith('http') or o == v else o)
            notes = [clean(g(r, c_note))]
            for extra in range(1, ws.max_column + 1):       # a note typed in a column with no heading
                if extra in H.values() or extra == 1: continue
                t = clean(g(r, extra))
                if t: notes.append(t)
            use_text = ' · '.join(n for n in notes if n)      # the NOTE is about the part's place on THIS sheet
            rmb, qq = num(g(r, c_rmb)), clean(g(r, c_qq))
            fields = {
                'itemCode': clean(g(r, c_id)), 'name': desc, 'color': clean(g(r, c_color)), 'size': clean(g(r, c_size)),
                'material': clean(g(r, c_mat)), 'weight': clean(g(r, c_wt)), 'vendor': vendor.replace('\n', ''), 'vendorSku': sku,
                'vendorUrl': url, 'origin': origin, 'priceUsd': num(g(r, c_usd)), 'dutyPct': num(g(r, c_duty)),
                'orderQty': clean(g(r, c_oq)), 'sourcingStatus': clean(g(r, c_status)), 'notes': '',
                'quotes': ([{'label': '', 'price': rmb, 'currency': 'RMB', 'qty': qq}] if rmb is not None else []),
            }
            key = (sku.upper() if sku else 'NAME:' + desc.upper())
            line = by_key.get(key)
            if line is not None and sec['key'] in line['uses']:
                # Two rows of ONE sheet are two parts, whatever the vendor number says — a number typed twice is
                # reported, never merged (Capa: 92406A247 sits on the 3/4" screw and on the 1" screw).
                report.append(f'"{ws.title}" rows {line["source"]["row"]} and {r} both carry vendor number {sku or "(none)"} — kept as two lines; check the number.')
                line = None
            if line is None:
                line = {'key': f'L{len(lines) + 1}', **fields, 'picture': '', 'uses': {}, 'source': {'sheet': ws.title, 'row': r}}
                by_key.setdefault(key, line)
                lines.append(line)
            else:
                for f in ('priceUsd', 'dutyPct'):
                    if fields[f] is not None and line[f] is not None and abs(fields[f] - line[f]) > 1e-9:
                        report.append(f'{sku or desc}: {f} is {line[f]} on "{line["source"]["sheet"]}" and {fields[f]} on "{ws.title}" row {r} — the first was kept.')
                for f, val in fields.items():
                    if line.get(f) in ('', None, []) and val not in ('', None, []): line[f] = val
            raw = in_cell.get(ws.title, {}).get(f'{openpyxl.utils.get_column_letter(c_pic)}{r}') if c_pic else None
            if raw and not line['picture']: line['picture'] = pics.add(raw, PART_PX, 'p')
            use_note = [use_text]
            if fields['size'] and fields['size'] != line['size']: use_note.insert(0, f'size here: {fields["size"]}')
            line['uses'][sec['key']] = {'qty': qty, 'balloon': balloon, 'note': ' · '.join(n for n in use_note if n)}
            each = num(g(r, c_each))
            if each is not None and fields['priceUsd'] is not None and fields['dutyPct'] is not None:
                if abs(each - fields['priceUsd'] * (1 + fields['dutyPct'])) > 0.005:
                    report.append(f'{sku or desc} on "{ws.title}" row {r}: the sheet says {each} each, USD × (1 + duty) is {round(fields["priceUsd"] * (1 + fields["dutyPct"]), 4)}.')
    return bundle(path, project_name, 'LIGHTING', suggested_brand, sections, lines, pics, report)


def hardware_quotes(path, project_name, suggested_brand):
    """H3: one item per pair of rows — a material on each row, a price at two quantities."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    pics, report, lines = Pictures(), [], []
    floats = floating_pictures(ws)
    H = header_map(ws, 3)
    c_note, c_id, c_name, c_mat, c_size = H.get('NOTES'), H.get('ITEM ID'), H.get('ITEM NAME'), H.get('INSTRUCTIONS'), H.get('SIZE')
    price_cols = [(ws.cell(3, c).column, clean(ws.cell(3, c).value)) for c in range(1, ws.max_column + 1) if clean(ws.cell(3, c).value).upper().startswith('PRICE')]
    line = None
    for r in range(4, ws.max_row + 1):
        code = clean(ws.cell(r, c_id).value)
        if code:
            line = {'key': f'L{len(lines) + 1}', 'itemCode': code, 'name': clean(ws.cell(r, c_name).value), 'size': clean(ws.cell(r, c_size).value),
                    'notes': clean(ws.cell(r, c_note).value) if c_note else '', 'quotes': [], 'picture': '', 'uses': {}, 'rows': [r], 'source': {'sheet': ws.title, 'row': r}}
            lines.append(line)
            if not re.match(r'^H3-', code): report.append(f'Row {r}: the item number reads "{code}" — the others on this sheet begin H3-. Brought in as typed.')
        elif line is not None:
            line['rows'].append(r)
        if line is None: continue
        mat = clean(ws.cell(r, c_mat).value)
        for (pc, label) in price_cols:
            raw_price = ws.cell(r, pc).value
            price = num(raw_price)
            if raw_price not in (None, '') and price is None: report.append(f'Row {r}: "{raw_price}" under {label} is not a number — left out.')
            if isinstance(raw_price, str) and '。' in raw_price: report.append(f'Row {r}: {label} was typed "{raw_price}" — read as {price}.')
            qty = clean(ws.cell(r, pc + 1).value) or next((q['qty'] for q in line['quotes'] if q['label2'] == label), '')
            if price is not None: line['quotes'].append({'label': mat, 'label2': label, 'price': price, 'currency': '', 'qty': qty})
    for line in lines:
        mine = [raw for (row, _c, raw) in floats if row in line['rows']]
        if mine: line['picture'] = pics.add(mine[0], PART_PX, 'p')
        if len(mine) > 1: report.append(f'{line["itemCode"]}: {len(mine)} pictures sit on its rows — the first was taken.')
        for q in line['quotes']: q['label'] = ' · '.join(x for x in (q.pop('label2').strip(), q['label']) if x)
        del line['rows']
    unplaced = len(floats) - sum(1 for l in lines if l['picture'])
    if unplaced > 0: report.append(f'{unplaced} picture(s) on the sheet sit on no item row and were not brought in.')
    report.append('Price 1 / Price 2 carry no currency on the sheet — brought in with the currency blank.')
    return bundle(path, project_name, 'HARDWARE', suggested_brand, [], lines, pics, report)


def bundle(path, project_name, kind, suggested_brand, sections, lines, pics, report):
    return {
        'format': 'control-sheet-bundle/1',
        'madeAt': datetime.datetime.now().isoformat(timespec='seconds'),
        'sourceFile': os.path.basename(path),
        'project': {'name': project_name, 'kind': kind, 'suggestedBrand': suggested_brand},
        'sections': sections, 'lines': lines, 'images': pics.images, 'report': report,
    }


JOBS = [
    ('M2CCapa Floor Lamp.xlsx', lighting, 'CAPA FLOOR LAMP', 'm2c'),
    ('M2C_Strata Light.xlsx', lighting, 'STRATA LIGHT', 'm2c'),
    ('H3-Contours.xlsx', hardware_quotes, 'H3 CONTOURS', 'ce'),
]

if __name__ == '__main__':
    os.makedirs(HERE, exist_ok=True)
    for (fname, fn, project, brand) in JOBS:
        src = os.path.join(SRC, fname)
        if not os.path.exists(src): print(f'— {fname}: not in Inception/, skipped'); continue
        b = fn(src, project, brand)
        out = os.path.join(HERE, re.sub(r'\.xlsx$', '', fname).replace(' ', '_') + '.control-sheet.json')
        with open(out, 'w') as f: json.dump(b, f)
        print(f'✓ {fname} → {os.path.basename(out)}  ({os.path.getsize(out) / 1e6:.1f} MB) · {len(b["sections"])} sheets · {len(b["lines"])} lines · {sum(1 for l in b["lines"] if l["picture"])} with a picture · {len(b["images"])} pictures')
        for s in b['sections']: print(f'    {s["key"]} {s["name"]} [{s["kind"]}] drawings {len(s["drawings"])} children {len(s["children"])} sheet total {s["sheetTotal"]}')
        for line in b['report']: print('    ! ' + line)
